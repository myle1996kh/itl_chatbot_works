#!/usr/bin/env python
"""
Ingest eTMS.docx into knowledge base for both tenants (eTMS and Google).
Uses DOCX extraction and PgVector for vector storage.
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from src.config import SessionLocal
from src.models.knowledge_document import KnowledgeDocument
from src.services.document_processor import DocumentProcessor
from src.services.embedding_service import EmbeddingService
from src.services.rag_service import RAGService
from src.utils.logging import get_logger
import uuid
from docx import Document as DocxDocument

logger = get_logger(__name__)

# Tenant IDs
TENANT_IDS = [
    "f160e26f-c41a-498f-9ab9-b3dbefbdbd50",  # eTMS
    "1193a40f-1d03-4ecd-a601-901a55589f56",  # Google
]

DOCX_PATH = "eTMS.docx"

def extract_text_from_docx(file_path: str) -> str:
    """Extract all text from DOCX file."""
    print(f"  Extracting text from {file_path}...")

    doc = DocxDocument(file_path)
    text_parts = []

    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text.strip())

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip():
                text_parts.append(row_text)

    full_text = "\n".join(text_parts)
    print(f"  Extracted {len(full_text)} characters, {len(text_parts)} sections")
    return full_text

def chunk_text(text: str, chunk_size: int = 600, overlap: int = 200) -> list:
    """Simple chunking strategy."""
    chunks = []
    step = chunk_size - overlap

    for i in range(0, len(text), step):
        chunk = text[i:i + chunk_size]
        if chunk.strip():
            chunks.append(chunk)

    return chunks

def ingest_for_tenant(db, tenant_id: str, text: str, chunks: list):
    """Ingest documents for a single tenant."""
    print(f"\n→ Ingesting for tenant: {tenant_id}")

    # Delete existing documents
    existing_count = db.query(KnowledgeDocument).filter(
        KnowledgeDocument.tenant_id == tenant_id
    ).count()

    if existing_count > 0:
        print(f"  Deleting {existing_count} existing documents...")
        db.query(KnowledgeDocument).filter(
            KnowledgeDocument.tenant_id == tenant_id
        ).delete()
        db.commit()

    # Create embeddings
    print(f"  Creating embeddings for {len(chunks)} chunks...")
    embedding_service = EmbeddingService()

    # Insert chunks
    inserted = 0
    for i, chunk in enumerate(chunks):
        if i % 100 == 0:
            print(f"    Processing chunk {i+1}/{len(chunks)}...")

        try:
            # Create embedding
            embedding = embedding_service.embed_text(chunk)

            # Create document record
            doc = KnowledgeDocument(
                document_id=uuid.uuid4(),
                tenant_id=tenant_id,
                filename="eTMS.docx",
                content=chunk,
                embedding=embedding,
                metadata={
                    "source": "eTMS.docx",
                    "chunk_index": i,
                    "total_chunks": len(chunks)
                }
            )

            db.add(doc)

            if (i + 1) % 50 == 0:
                db.commit()
                print(f"    Committed {i+1} documents...")

            inserted += 1

        except Exception as e:
            logger.error("chunk_ingestion_failed", chunk_index=i, error=str(e))
            db.rollback()
            continue

    # Final commit
    db.commit()
    print(f"  ✅ Inserted {inserted} documents for {tenant_id}")

def main():
    """Main ingestion flow."""
    print("\n" + "="*70)
    print("DOCUMENT INGESTION: eTMS.docx")
    print("="*70)

    db = SessionLocal()

    try:
        # Check if DOCX file exists
        if not Path(DOCX_PATH).exists():
            print(f"❌ File not found: {DOCX_PATH}")
            return False

        # Extract text
        print(f"\n→ Processing {DOCX_PATH}...")
        full_text = extract_text_from_docx(DOCX_PATH)

        # Chunk text
        print(f"\n→ Creating chunks (600 chars, 200 overlap)...")
        chunks = chunk_text(full_text, chunk_size=600, overlap=200)
        print(f"  Created {len(chunks)} chunks")

        # Ingest for both tenants
        for tenant_id in TENANT_IDS:
            ingest_for_tenant(db, tenant_id, full_text, chunks)

        print("\n" + "="*70)
        print("✅ Document ingestion completed successfully!")
        print("="*70 + "\n")
        return True

    except Exception as e:
        print(f"\n❌ Ingestion failed: {str(e)}")
        logger.error("ingestion_failed", error=str(e))
        db.rollback()
        return False

    finally:
        db.close()

if __name__ == "__main__":
    success = main()
    sys.exit(0 if success else 1)
