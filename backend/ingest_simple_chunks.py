"""
Script to ingest documents using simple character-based chunking (like sample_rag.md)
- Deletes existing documents for the tenant
- Chunks text by character count (400 characters per chunk)
- Extracts section titles from headings
- Saves to PostgreSQL with pgvector using LangChain schema
"""

import psycopg2
from psycopg2.extras import execute_batch, RealDictCursor
from docx import Document
from sentence_transformers import SentenceTransformer
import uuid
import json

# =============================
# CONFIGURATION
# =============================
TENANT_ID = "1193a40f-1d03-4ecd-a601-901a55589f56"
DOCUMENT_PATH = "eTMS.docx"
CHUNK_SIZE = 400  # characters per chunk
BATCH_SIZE = 64
COLLECTION_NAME = "knowledge_documents"  # Collection name for this tenant

# Database connection
DB_CONFIG = {
    "host": "localhost",
    "port": 5432,
    "database": "chatbot_itl",
    "user": "postgres",
    "password": "123456"
}

print("=" * 100)
print("SIMPLE CHARACTER-BASED DOCUMENT INGESTION")
print("=" * 100)
print(f"Tenant ID: {TENANT_ID}")
print(f"Document: {DOCUMENT_PATH}")
print(f"Chunk size: {CHUNK_SIZE} characters")
print(f"Batch size: {BATCH_SIZE}")
print(f"Collection: {COLLECTION_NAME}")

# =============================
# 1️⃣ CONNECT TO DATABASE
# =============================
print("\n1. CONNECTING TO DATABASE")
print("-" * 100)

conn = psycopg2.connect(**DB_CONFIG)
cur = conn.cursor(cursor_factory=RealDictCursor)

print("✅ Connected to PostgreSQL")

# =============================
# 2️⃣ GET OR CREATE COLLECTION
# =============================
print("\n2. MANAGING COLLECTION")
print("-" * 100)

# Check if collection exists
cur.execute("""
    SELECT uuid, name FROM langchain_pg_collection WHERE name = %s
""", (COLLECTION_NAME,))

collection = cur.fetchone()

if collection:
    collection_id = collection['uuid']
    print(f"✅ Found existing collection: {COLLECTION_NAME}")
    print(f"   Collection ID: {collection_id}")

    # Delete existing embeddings for this tenant
    cur.execute("""
        DELETE FROM langchain_pg_embedding
        WHERE collection_id = %s
        AND cmetadata->>'tenant_id' = %s
    """, (collection_id, TENANT_ID))

    deleted_count = cur.rowcount
    conn.commit()
    print(f"✅ Deleted {deleted_count} existing embeddings for tenant")
else:
    # Create new collection
    collection_id = uuid.uuid4()
    cur.execute("""
        INSERT INTO langchain_pg_collection (uuid, name, cmetadata)
        VALUES (%s, %s, %s)
        RETURNING uuid, name
    """, (collection_id, COLLECTION_NAME, json.dumps({})))

    collection = cur.fetchone()
    conn.commit()
    print(f"✅ Created new collection: {COLLECTION_NAME}")
    print(f"   Collection ID: {collection['uuid']}")

# =============================
# 3️⃣ LOAD DOCX FILE
# =============================
print("\n3. LOADING DOCX FILE")
print("-" * 100)

doc = Document(DOCUMENT_PATH)
paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

print(f"✅ Loaded {len(paragraphs)} paragraphs from {DOCUMENT_PATH}")

# =============================
# 4️⃣ EXTRACT HEADINGS (SECTION TITLES)
# =============================
print("\n4. EXTRACTING SECTION TITLES")
print("-" * 100)

headings = []
for para in doc.paragraphs:
    if para.style.name.startswith("Heading"):
        headings.append(para.text.strip())

print(f"✅ Found {len(headings)} headings:")
for i, heading in enumerate(headings[:10], 1):
    print(f"  {i}. {heading}")
if len(headings) > 10:
    print(f"  ... and {len(headings) - 10} more")

# =============================
# 5️⃣ CHUNK TEXT BY CHARACTER COUNT
# =============================
print(f"\n5. CHUNKING TEXT (SIZE={CHUNK_SIZE} characters)")
print("-" * 100)

def chunk_text_by_chars(text, size=400):
    """Split text into chunks by character count"""
    chunks = []
    for i in range(0, len(text), size):
        chunk = text[i:i+size]
        if chunk.strip():  # Only add non-empty chunks
            chunks.append(chunk)
    return chunks

chunks = []
meta_info = []  # (page_number, section_title)

current_section = "Unknown"
for i, paragraph in enumerate(paragraphs):
    # Check if this paragraph is a heading
    for heading in headings:
        if heading == paragraph or heading in paragraph:
            current_section = heading
            break

    sub_chunks = chunk_text_by_chars(paragraph, CHUNK_SIZE)
    for chunk in sub_chunks:
        page_number = i // 10 + 1  # Assume 10 paragraphs ~ 1 page

        chunks.append(chunk)
        meta_info.append((page_number, current_section))

print(f"✅ Created {len(chunks)} chunks")
print(f"   Average chunk size: {sum(len(c) for c in chunks) // len(chunks)} characters")

# Show section distribution
section_counts = {}
for _, section in meta_info:
    section_counts[section] = section_counts.get(section, 0) + 1

print(f"\n📊 Chunks by section (top 10):")
for section, count in sorted(section_counts.items(), key=lambda x: x[1], reverse=True)[:10]:
    print(f"   {count:4d} chunks - {section[:80]}")

# =============================
# 6️⃣ GENERATE EMBEDDINGS AND SAVE
# =============================
print("\n6. GENERATING EMBEDDINGS AND SAVING TO DATABASE")
print("-" * 100)

embedder = SentenceTransformer("all-MiniLM-L6-v2")
print("✅ Loaded embedding model: all-MiniLM-L6-v2 (384 dimensions)")

total_batches = (len(chunks) + BATCH_SIZE - 1) // BATCH_SIZE

for i in range(0, len(chunks), BATCH_SIZE):
    batch = chunks[i:i+BATCH_SIZE]
    batch_meta = meta_info[i:i+BATCH_SIZE]

    # Generate embeddings
    embeddings = embedder.encode(batch)

    # Prepare data for insertion
    data = []
    for j in range(len(batch)):
        embedding_id = str(uuid.uuid4())
        content = batch[j]
        embedding = embeddings[j].tolist()
        page_number, section_title = batch_meta[j]

        # Calculate chunk index
        chunk_index = i + j

        # Create metadata
        metadata = {
            "tenant_id": TENANT_ID,
            "document_name": DOCUMENT_PATH,
            "page_number": page_number,
            "section_title": section_title,
            "chunk_index": chunk_index,
            "chunk_size": len(content)
        }

        data.append((
            embedding_id,
            collection_id,
            embedding,
            content,
            json.dumps(metadata)
        ))

    # Batch insert
    execute_batch(
        cur,
        """
        INSERT INTO langchain_pg_embedding (
            id,
            collection_id,
            embedding,
            document,
            cmetadata
        ) VALUES (%s, %s, %s::vector, %s, %s::jsonb)
        """,
        data
    )
    conn.commit()

    batch_num = i // BATCH_SIZE + 1
    print(f"✅ Saved batch {batch_num}/{total_batches} ({len(batch)} chunks)")

# =============================
# 7️⃣ VERIFY INGESTION
# =============================
print("\n7. VERIFYING INGESTION")
print("-" * 100)

# Count total chunks
cur.execute("""
    SELECT COUNT(*) as total_chunks
    FROM langchain_pg_embedding
    WHERE collection_id = %s
    AND cmetadata->>'tenant_id' = %s
""", (collection_id, TENANT_ID))
result = cur.fetchone()
print(f"✅ Total chunks in database: {result['total_chunks']}")

# Get statistics by section
cur.execute("""
    SELECT
        cmetadata->>'section_title' as section_title,
        COUNT(*) as chunk_count,
        AVG(LENGTH(document)) as avg_length
    FROM langchain_pg_embedding
    WHERE collection_id = %s
    AND cmetadata->>'tenant_id' = %s
    GROUP BY cmetadata->>'section_title'
    ORDER BY chunk_count DESC
    LIMIT 10
""", (collection_id, TENANT_ID))

print("\n📊 Top 10 sections by chunk count:")
for row in cur.fetchall():
    print(f"   {row['chunk_count']:4d} chunks, avg {int(row['avg_length']):3d} chars - {row['section_title'][:70]}")

# Sample chunks
cur.execute("""
    SELECT
        LEFT(document, 100) as preview,
        LENGTH(document) as length,
        cmetadata->>'section_title' as section_title
    FROM langchain_pg_embedding
    WHERE collection_id = %s
    AND cmetadata->>'tenant_id' = %s
    ORDER BY (cmetadata->>'chunk_index')::int
    LIMIT 5
""", (collection_id, TENANT_ID))

print("\n📝 Sample chunks:")
for chunk in cur.fetchall():
    print(f"   [{chunk['length']} chars] {chunk['section_title'][:50]}")
    print(f"   {chunk['preview']}...")

# Close connection
cur.close()
conn.close()

print("\n" + "=" * 100)
print("✅ INGESTION COMPLETE")
print("=" * 100)
print(f"\nNext steps:")
print(f"1. Test RAG query with tenant: {TENANT_ID}")
print(f"2. Compare results with previous chunk_size=600 approach")
print(f"""
Test query example:
curl -X POST 'http://localhost:8000/api/{TENANT_ID}/chat' \\
  -H 'Content-Type: application/json' \\
  -d '{{"message": "Hướng dẫn tạo mới bảng báo giá bán LCL?", "user_id": "test_user"}}'
""")
